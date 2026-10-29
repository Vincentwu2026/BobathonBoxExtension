#!/usr/bin/env python3
"""
Generates the 8 demo documents in ./demo-files (already committed; only re-run to change them).

Files are written by real Office libraries, so the preview code is exercised on genuine output.
All names, emails and figures are fictional.

Usage:  python3 scripts/make-demo-files.py demo-files
Needs:  pip install openpyxl python-pptx python-docx   (use a virtualenv)
"""

import json
import sys
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font
from pptx import Presentation


def text(t, marks=None):
    node = {"type": "text", "text": t}
    if marks:
        node["marks"] = [{"type": m} for m in marks]
    return node


def para(*nodes):
    return {"type": "paragraph", "content": list(nodes)}


def bullets(items):
    return {"type": "bullet_list", "content": [{"type": "list_item", "content": [para(text(i))]} for i in items]}


def heading(level, t):
    return {"type": "heading", "attrs": {"level": level}, "content": [text(t)]}


def make_boxnote(path: Path) -> None:
    doc = {
        "type": "doc",
        "content": [
            heading(1, "Q4 Planning Meeting"),
            para(text("Attendees: Priya Nair, Marcus Webb, Elena Rossi.")),
            heading(2, "Action items"),
            bullets([
                "Priya: publish the revised revenue forecast by Friday",
                "Marcus: confirm the Northwind renewal date",
                "Elena: finish the vendor contract review",
            ]),
            para(text("Next meeting: first Monday of November.", ["strong"])),
        ],
    }
    path.write_text(json.dumps({"version": 1, "schema_version": 1, "doc": doc}), encoding="utf-8")


def make_account_plan(path: Path) -> None:
    d = Document()
    d.add_heading("Northwind Bank Account Plan", level=1)
    d.add_paragraph("Prepared by the enterprise accounts team.")
    d.add_heading("Objectives", level=2)
    for item in ("Expand the hybrid cloud footprint", "Renew mainframe support through FY28"):
        d.add_paragraph(item, style="List Bullet")
    p = d.add_paragraph("Total contract value: ")
    p.add_run("$4.2M").bold = True
    t = d.add_table(rows=3, cols=3)
    for r, row in enumerate([("Quarter", "Target", "Status"), ("Q3", "$1.1M", "Closed"), ("Q4", "$1.4M", "At risk")]):
        for c, value in enumerate(row):
            t.cell(r, c).text = value
    d.save(path)


def make_contract_summary(path: Path) -> None:
    d = Document()
    d.add_heading("Vendor Contract Summary", level=1)
    d.add_paragraph("Supplier: Contoso Cloud Services. Term: 3 years from January 2026.")
    d.add_heading("Key terms", level=2)
    for item in ("Annual fee: $240,000, paid quarterly", "Either side may end the contract with 90 days notice", "Data stays in the EU"):
        d.add_paragraph(item, style="List Bullet")
    d.add_paragraph("Owner: Elena Rossi. Legal review completed in September.")
    d.save(path)


def make_roadmap(path: Path) -> None:
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Product Roadmap FY27"
    s1.placeholders[1].text = "Product team"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Three priorities"
    body = s2.placeholders[1].text_frame
    body.text = "Faster search across all files"
    body.add_paragraph().text = "Safer deletion with retention checks"
    body.add_paragraph().text = "AI answers about any document"
    s2.notes_slide.notes_text_frame.text = "Lead with search; it is the most requested feature."
    prs.save(path)


def make_training(path: Path) -> None:
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Security Awareness Training"
    s1.placeholders[1].text = "Annual refresher"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Five simple rules"
    body = s2.placeholders[1].text_frame
    body.text = "Use a password manager"
    for line in ("Turn on multi-factor sign-in", "Check the sender before you click", "Lock your screen when you leave", "Report anything odd straight away"):
        body.add_paragraph().text = line
    prs.save(path)


def make_forecast(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Forecast"
    ws.append(["Region", "Q3 actual ($M)", "Q4 forecast ($M)"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in (("EMEA", 118.4, 128.0), ("Americas", 201.3, 233.5), ("APAC", 74.3, 81.7)):
        ws.append(list(row))
    wb.save(path)


def make_survey(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Responses"
    ws.append(["Customer", "Score (1-10)", "Comment"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in (
        ("Northwind Bank", 9, "Support is quick and friendly"),
        ("Fabrikam Retail", 7, "Reports are slow to load"),
        ("Adventure Works", 8, "Great onboarding"),
        ("Tailspin Travel", 5, "Pricing is unclear"),
    ):
        ws.append(list(row))
    wb.save(path)


def main() -> None:
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    out = Path(sys.argv[1])
    out.mkdir(parents=True, exist_ok=True)
    make_boxnote(out / "Q4_Planning_Meeting.boxnote")
    make_account_plan(out / "Northwind_Account_Plan.docx")
    make_contract_summary(out / "Vendor_Contract_Summary.docx")
    make_roadmap(out / "Product_Roadmap.pptx")
    make_training(out / "Security_Training.pptx")
    make_forecast(out / "Q4_Revenue_Forecast.xlsx")
    make_survey(out / "Customer_Feedback_Survey.xlsx")
    (out / "Team_Offsite_Ideas.txt").write_text(
        "Team offsite ideas\n\n- Cooking class\n- Escape room\n- Walk and picnic by the river\n\nVote by the end of the month.\n",
        encoding="utf-8",
    )
    # What a download cannot carry: who owns each file, who it is shared with, and do-not-delete flags.
    governance = {
        "Q4_Planning_Meeting.boxnote": {"owner": "priya.nair@example.com", "collaborators": ["marcus.webb@example.com", "elena.rossi@example.com"]},
        "Northwind_Account_Plan.docx": {"owner": "priya.nair@example.com", "collaborators": ["marcus.webb@example.com"]},
        "Q4_Revenue_Forecast.xlsx": {"owner": "elena.rossi@example.com", "collaborators": ["priya.nair@example.com"]},
        "Product_Roadmap.pptx": {"owner": "marcus.webb@example.com", "doNotDelete": True},
        "Vendor_Contract_Summary.docx": {"owner": "elena.rossi@example.com", "doNotDelete": True},
        "Customer_Feedback_Survey.xlsx": {"owner": "marcus.webb@example.com"},
        "Security_Training.pptx": {"owner": "elena.rossi@example.com"},
        "Team_Offsite_Ideas.txt": {"owner": "priya.nair@example.com"},
    }
    (out / "governance.json").write_text(json.dumps(governance, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote demo files to {out}")


if __name__ == "__main__":
    main()
